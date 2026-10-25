"""Document-integrity regression checks for the built manuscript.docx.

These tests assert structural/editorial invariants of the generated DOCX —
not the scientific values themselves (those are covered by numeric-provenance
and p12_qc). Run after build_manuscript.py.
"""
import re
import zipfile
from pathlib import Path

import docx
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "submission" / "manuscript.docx"


def _doc():
    return docx.Document(str(DOCX))


def _xml():
    with zipfile.ZipFile(DOCX) as z:
        return z.read("word/document.xml").decode("utf8")


def _body_text(d):
    return "\n".join(p.text for p in d.paragraphs)


def test_omml_equations_present_and_editable():
    x = _xml()
    # display equations for WFI, per-event WERR and WERR_CD
    assert x.count("<m:oMath") >= 3
    assert x.count("<m:nary") >= 2  # two summations (WFI, WERR_CD numerator/denominator)
    assert x.count("<m:f>") >= 2   # genuine fractions
    # no image-based equations: math regions must not contain drawings
    for m in re.finditer(r"<m:oMathPara.*?</m:oMathPara>", x, re.S):
        assert "<w:drawing>" not in m.group(0) and "<w:pict" not in m.group(0)


def test_figure_caption_style_and_order():
    d = _doc()
    caps = [p.text for p in d.paragraphs if re.match(r"^Fig\. \d+\.", p.text.strip())]
    assert len(caps) == 8
    nums = [int(re.match(r"^Fig\. (\d+)\.", c.strip()).group(1)) for c in caps]
    assert nums == list(range(1, 9))
    # no legacy "Figure N." captions
    assert not any(re.match(r"^Figure \d+\.", p.text.strip()) for p in d.paragraphs)


def test_first_mention_order_figures_and_tables():
    d = _doc()
    caps_start = next(i for i, p in enumerate(d.paragraphs) if re.match(r"^Fig\. 1\.", p.text.strip()))
    body = d.paragraphs[:caps_start]
    fig_seq = [int(m.group(1)) for p in body for m in re.finditer(r"Fig\. (\d+)", p.text)]
    tab_seq = [int(m.group(1)) for p in body for m in re.finditer(r"Table (\d+)", p.text)]
    first_fig = sorted(set(fig_seq), key=fig_seq.index)
    first_tab = sorted(set(tab_seq), key=tab_seq.index)
    assert first_fig == sorted(first_fig)
    assert first_tab == sorted(first_tab)


def test_no_supplement_references_without_supplement():
    text = _body_text(_doc())
    assert not re.search(r"Fig\. S\d", text)
    assert not re.search(r"Table S\d", text)


def test_no_markdown_asterisk_remnants():
    x = _xml()
    # literal '*' inside text runs
    for m in re.finditer(r"<w:t[^>]*>([^<]*)</w:t>", x):
        assert "*" not in m.group(1)


def test_no_character_count_grid_or_eastasian_snap():
    x = _xml()
    assert 'w:type="linesAndChars"' not in x
    assert 'w:type="snapToChars"' not in x
    assert "snapToGrid" not in x


def test_werr_cd_matches_canonical_csv():
    canon = pd.read_csv(ROOT / "results" / "country_day_werr_summary.csv").set_index("country").loc["ALL"]
    text = _body_text(_doc())
    assert f"{canon.werr_cd:+.4f}" in text or f"{canon.werr_cd:.4f}" in text or f"{canon.werr_cd:.3f}" in text
    lo = f"{canon.boot_lo:.4f}"; hi = f"{canon.boot_hi:.4f}"
    assert lo in text and hi in text


def test_pooled_event_mean_distinct_from_werr_cd():
    pooled = pd.read_csv(ROOT / "results" / "werr_pooled_summary.csv").iloc[0]
    text = _body_text(_doc())
    assert f"{pooled.werr_mean:.3f}" in text  # -0.036 event mean, separate estimand
