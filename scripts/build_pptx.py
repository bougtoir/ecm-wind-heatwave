from pptx import Presentation
from pptx.util import Inches, Pt
import pathlib
F = pathlib.Path("figs")
prs = Presentation(); prs.slide_width = Inches(10); prs.slide_height = Inches(7.5)
caps = {
 "fig1_deployment": "Fig.1 Deployment & transects",
 "fig2_moisture": "Fig.2 MFC climatology & transect fluxes",
 "fig3_heatwaves": "Fig.3 Heatwave climatology",
 "fig4_panel": "Fig.4 Regional panel coefficients",
 "fig5_analogs": "Fig.5 Analog-matched effects",
 "fig6_eventstudy": "Fig.6 Deployment event study",
 "fig7_falsification": "Fig.7 Falsification battery",
 "fig8_werr": "Fig.8 WERR per country",
}
for name, cap in caps.items():
    s = prs.slides.add_slide(prs.slide_layouts[5])
    tb = s.shapes.add_textbox(Inches(0.3), Inches(0.2), Inches(9), Inches(0.5))
    tb.text_frame.text = cap; tb.text_frame.paragraphs[0].font.size = Pt(20)
    s.shapes.add_picture(str(F/f"{name}.png"), Inches(0.5), Inches(0.9),
                         width=Inches(9))
prs.save("manuscript/figures_editable.pptx")
print("pptx saved")
